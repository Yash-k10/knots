"""
KNOTS Platform - Comprehensive System Architecture & Developer Manual Generator
Generates a complete, professional Microsoft Word (.docx) document containing
every detail, algorithm, code reference, API catalog, database schema, and guide.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background color for a table cell."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    """Sets internal padding for a cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Sets clean borders for an entire table."""
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="none"/>
        <w:right w:val="none"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def add_callout(doc, text, title="NOTE", box_type="info"):
    """Adds a stylish callout box using a single-cell table with colored left border."""
    colors = {
        "info": {"bg": "F0F9FF", "border": "0284C7", "title_color": RGBColor(2, 132, 199)},
        "tip": {"bg": "F0FDF4", "border": "16A34A", "title_color": RGBColor(22, 163, 74)},
        "warning": {"bg": "FFFBEB", "border": "D97706", "title_color": RGBColor(217, 119, 6)},
        "important": {"bg": "EEF2FF", "border": "4F46E5", "title_color": RGBColor(79, 70, 229)}
    }
    cfg = colors.get(box_type, colors["info"])
    
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, cfg["bg"])
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Custom border with thick left line
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'''
    <w:tcBorders {nsdecls("w")}>
        <w:left w:val="single" w:sz="24" w:space="0" w:color="{cfg['border']}"/>
        <w:top w:val="none"/>
        <w:right w:val="none"/>
        <w:bottom w:val="none"/>
    </w:tcBorders>
    '''
    tcPr.append(parse_xml(borders_xml))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"[{title}] ")
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(10)
    run_title.font.color.rgb = cfg["title_color"]
    
    run_text = p.add_run(text)
    run_text.font.name = "Arial"
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor(51, 65, 85)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_table(table, col_widths, col_alignments=None):
    """Applies executive styling to standard tables."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    
    # Header Row
    header_row = table.rows[0]
    header_tr = header_row._tr.get_or_add_trPr()
    header_tr.append(OxmlElement('w:tblHeader'))
    
    for i, cell in enumerate(header_row.cells):
        cell.width = col_widths[i]
        set_cell_background(cell, "1E3A8A") # Navy
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        if col_alignments and i < len(col_alignments):
            p.alignment = col_alignments[i]
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            
    # Body Rows
    for row_idx, row in enumerate(table.rows[1:], start=1):
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for i, cell in enumerate(row.cells):
            cell.width = col_widths[i]
            if bg_color != "FFFFFF":
                set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if col_alignments and i < len(col_alignments):
                p.alignment = col_alignments[i]
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(30, 41, 59)

def add_code_block(doc, code_str):
    """Adds a formatted monospace code block."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "1E293B") # Dark slate
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_str)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(241, 245, 249) # Off-white
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

print("Setup helper functions ready.")
