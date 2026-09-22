import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_demo_doc():
    doc = Document()

    # Set Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Color definitions
    PRIMARY = RGBColor(75, 99, 210)       # #4B63D2
    DARK = RGBColor(30, 39, 70)          # #1E2746
    SLATE = RGBColor(88, 81, 164)        # #5851A4
    EMERALD = RGBColor(16, 122, 87)      # #107A57

    def style_header_cell(cell, text, width_in=None):
        cell.text = text
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="4B63D2"/>')
        cell._tc.get_or_add_tcPr().append(shd)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)
            run.font.name = 'Segoe UI'
        if width_in:
            cell.width = Inches(width_in)

    def style_row_cell(cell, text, is_even=False, width_in=None, is_bold=False, is_code=False):
        cell.text = text
        if is_even:
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F4F2FA"/>')
            cell._tc.get_or_add_tcPr().append(shd)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = is_bold
            run.font.size = Pt(9)
            run.font.name = 'Consolas' if is_code else 'Segoe UI'
            run.font.color.rgb = DARK if not is_code else PRIMARY
        if width_in:
            cell.width = Inches(width_in)

    # Title Header
    p_title = doc.add_paragraph()
    r_title = p_title.add_run('KNOTS Platform — Demo Accounts & Credentials')
    r_title.font.name = 'Segoe UI'
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = PRIMARY
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_after = Pt(2)

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run('Official Reference Document for Project Viva, Demonstration & Feature Evaluation')
    r_sub.font.name = 'Segoe UI'
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = SLATE
    p_sub.paragraph_format.space_after = Pt(12)

    # Universal Password Callout Box
    tbl_box = doc.add_table(rows=1, cols=1)
    tbl_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_box = tbl_box.rows[0].cells[0]
    shd_box = parse_xml(f'<w:shd {nsdecls("w")} w:fill="EBF0FF"/>')
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="4B63D2"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    c_box._tc.get_or_add_tcPr().append(shd_box)
    c_box._tc.get_or_add_tcPr().append(borders)
    p_box = c_box.paragraphs[0]
    r_box1 = p_box.add_run('🔑 Universal Default Password: ')
    r_box1.bold = True
    r_box1.font.color.rgb = DARK
    r_box2 = p_box.add_run('pass@knots')
    r_box2.bold = True
    r_box2.font.color.rgb = PRIMARY
    r_box2.font.size = Pt(12)
    r_box3 = p_box.add_run('\nAll pre-seeded demo accounts across every role share this standard password. All accounts are active and email-verified in the database.')
    r_box3.font.color.rgb = DARK
    r_box3.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section Table Generator
    def create_section_table(title, subtitle, headers, data, widths):
        p_h = doc.add_paragraph()
        r_h = p_h.add_run(title)
        r_h.font.name = 'Segoe UI'
        r_h.font.size = Pt(13)
        r_h.font.bold = True
        r_h.font.color.rgb = PRIMARY
        p_h.paragraph_format.space_before = Pt(10)
        p_h.paragraph_format.space_after = Pt(2)
        
        if subtitle:
            p_sub = doc.add_paragraph()
            r_s = p_sub.add_run(subtitle)
            r_s.font.size = Pt(9.5)
            r_s.font.italic = True
            r_s.font.color.rgb = SLATE
            p_sub.paragraph_format.space_after = Pt(6)

        table = doc.add_table(rows=len(data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Header Row
        for col_idx, (h_text, w) in enumerate(zip(headers, widths)):
            style_header_cell(table.rows[0].cells[col_idx], h_text, w)

        # Data Rows
        for row_idx, row_data in enumerate(data):
            is_even = (row_idx % 2 == 1)
            for col_idx, (val, w) in enumerate(zip(row_data, widths)):
                cell = table.rows[row_idx + 1].cells[col_idx]
                is_code = (col_idx == 1 or col_idx == 2)
                is_bold = (col_idx == 0)
                style_row_cell(cell, val, is_even=is_even, width_in=w, is_bold=is_bold, is_code=is_code)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(10)

    headers = ['Role', 'Email Address', 'Password', 'Name / Details']
    widths = [1.3, 2.3, 1.1, 2.3]

    # 1. Leadership & Central Admin
    data_leadership = [
        ['Central Admin', 'centraladmin@sbjit.edu.in', 'pass@knots', 'Central Administrator (Master Admin Access)'],
        ['TPO', 'tpo@sbjit.edu.in', 'pass@knots', 'Training & Placement Officer (Job & Candidate Hub)'],
    ]
    create_section_table('1. Campus Leadership & Central Administration', 'Institutional level master administrators', headers, data_leadership, widths)

    # 2. AIML Department
    data_aiml = [
        ['HOD (AIML)', 'hod.aiml@sbjit.edu.in', 'pass@knots', 'Dr. Manisha Kulkarni (AIML HOD Dashboard)'],
        ['Controller (AIML)', 'controller.aiml@sbjit.edu.in', 'pass@knots', 'Prof. Deepak Verma (Dept Approvals & Clubs)'],
        ['Faculty 1 (AIML)', 'faculty1.aiml@sbjit.edu.in', 'pass@knots', 'Dr. Pradeep Mishra (Poster: ML & AI Opportunities)'],
        ['Faculty 2 (AIML)', 'faculty2.aiml@sbjit.edu.in', 'pass@knots', 'Prof. Kavita Rao (AI Researcher)'],
        ['Faculty 3 (AIML)', 'faculty3.aiml@sbjit.edu.in', 'pass@knots', 'Dr. Sanjay Trivedi (Data Science Lab Lead)'],
        ['Faculty 4 (AIML)', 'faculty4.aiml@sbjit.edu.in', 'pass@knots', 'Prof. Anjali Somani (NLP & Vision Projects)'],
        ['Student 1 (AIML)', 'student1.aiml@sbjit.edu.in', 'pass@knots', 'Aryan Kapoor (AIML 4th Year)'],
        ['Student 2 (AIML)', 'student2.aiml@sbjit.edu.in', 'pass@knots', 'Isha Sen (AIML 4th Year)'],
        ['Student 3 (AIML)', 'student3.aiml@sbjit.edu.in', 'pass@knots', 'Varun Malhotra (AIML 4th Year)'],
        ['Student 4 (AIML)', 'student4.aiml@sbjit.edu.in', 'pass@knots', 'Mehak Chawla (AIML 3rd Year)'],
        ['Student 5 (AIML)', 'student5.aiml@sbjit.edu.in', 'pass@knots', 'Dev Singhania (AIML 3rd Year)'],
        ['Student 6 (AIML)', 'student6.aiml@sbjit.edu.in', 'pass@knots', 'Shreya Ghoshal (AIML 3rd Year)'],
        ['Student 7 (AIML)', 'student7.aiml@sbjit.edu.in', 'pass@knots', 'Kabir Bedi (AIML 2nd Year)'],
        ['Student 8 (AIML)', 'student8.aiml@sbjit.edu.in', 'pass@knots', 'Simran Kaur (AIML 2nd Year)'],
        ['Student 9 (AIML)', 'student9.aiml@sbjit.edu.in', 'pass@knots', 'Ayush Roy (AIML 1st Year)'],
        ['Student 10 (AIML)', 'student10.aiml@sbjit.edu.in', 'pass@knots', 'Kriti Sanon (AIML 1st Year)'],
        ['Alumni 1 (AIML)', 'alumni1.aiml@sbjit.edu.in', 'pass@knots', 'Sameer Sheikh (AI Engineer @ Microsoft)'],
        ['Alumni 2 (AIML)', 'alumni2.aiml@sbjit.edu.in', 'pass@knots', 'Nidhi Agrawal (Data Scientist @ Amazon)'],
        ['Alumni 3 (AIML)', 'alumni3.aiml@sbjit.edu.in', 'pass@knots', 'Gaurav Taneja (MLOps @ Google)'],
        ['Alumni 4 (AIML)', 'alumni4.aiml@sbjit.edu.in', 'pass@knots', 'Pallavi Shrestha (Research Scientist @ Adobe)'],
        ['Alumni 5 (AIML)', 'alumni5.aiml@sbjit.edu.in', 'pass@knots', 'Nikhil Kamath (Founder / AI Architect)'],
    ]
    create_section_table('2. Department of Artificial Intelligence & Machine Learning (AIML)', 'Faculty, HOD, Controller, Students (Years 1-4) & Verified Alumni', headers, data_aiml, widths)

    # 3. CSE Department
    data_cse = [
        ['HOD (CSE)', 'hod.cse@sbjit.edu.in', 'pass@knots', 'Dr. Arvind Sharma (CSE HOD Dashboard)'],
        ['Controller (CSE)', 'controller.cse@sbjit.edu.in', 'pass@knots', 'Prof. Amit Saxena (Dept Approvals & Clubs)'],
        ['Faculty 1 (CSE)', 'faculty1.cse@sbjit.edu.in', 'pass@knots', 'Dr. Rajesh Kumar (Distributed Systems)'],
        ['Faculty 2 (CSE)', 'faculty2.cse@sbjit.edu.in', 'pass@knots', 'Prof. Sunita Rao (Software Architecture)'],
        ['Faculty 3 (CSE)', 'faculty3.cse@sbjit.edu.in', 'pass@knots', 'Dr. Vikram Patil (Cloud Infrastructure)'],
        ['Faculty 4 (CSE)', 'faculty4.cse@sbjit.edu.in', 'pass@knots', 'Prof. Neha Deshpande (Full-Stack Mentor)'],
        ['Student 1 (CSE)', 'student1.cse@sbjit.edu.in', 'pass@knots', 'Aarav Sharma (CSE 4th Year)'],
        ['Student 2 (CSE)', 'student2.cse@sbjit.edu.in', 'pass@knots', 'Priya Patel (CSE 4th Year)'],
        ['Student 3 (CSE)', 'student3.cse@sbjit.edu.in', 'pass@knots', 'Rohan Verma (CSE 4th Year)'],
        ['Student 4 (CSE)', 'student4.cse@sbjit.edu.in', 'pass@knots', 'Ananya Iyer (CSE 3rd Year)'],
        ['Student 5 (CSE)', 'student5.cse@sbjit.edu.in', 'pass@knots', 'Aditya Joshi (CSE 3rd Year)'],
        ['Student 6 (CSE)', 'student6.cse@sbjit.edu.in', 'pass@knots', 'Sneha Kulkarni (CSE 3rd Year)'],
        ['Student 7 (CSE)', 'student7.cse@sbjit.edu.in', 'pass@knots', 'Tanmay Deshmukh (CSE 2nd Year)'],
        ['Student 8 (CSE)', 'student8.cse@sbjit.edu.in', 'pass@knots', 'Riya Gupta (CSE 2nd Year)'],
        ['Student 9 (CSE)', 'student9.cse@sbjit.edu.in', 'pass@knots', 'Harsh Mehta (CSE 1st Year)'],
        ['Student 10 (CSE)', 'student10.cse@sbjit.edu.in', 'pass@knots', 'Pooja Nair (CSE 1st Year)'],
        ['Alumni 1 (CSE)', 'alumni1.cse@sbjit.edu.in', 'pass@knots', 'Kunal Shinde (SDE @ Amazon)'],
        ['Alumni 2 (CSE)', 'alumni2.cse@sbjit.edu.in', 'pass@knots', 'Meera Bhatt (Backend Engineer @ Uber)'],
        ['Alumni 3 (CSE)', 'alumni3.cse@sbjit.edu.in', 'pass@knots', 'Siddharth Jain (Full-Stack Lead @ Atlassian)'],
        ['Alumni 4 (CSE)', 'alumni4.cse@sbjit.edu.in', 'pass@knots', 'Divya Wagh (Security Analyst @ Cisco)'],
        ['Alumni 5 (CSE)', 'alumni5.cse@sbjit.edu.in', 'pass@knots', 'Akash Mohite (Cloud Solutions @ Oracle)'],
    ]
    create_section_table('3. Department of Computer Science & Engineering (CSE)', 'Faculty, HOD, Controller, Students (Years 1-4) & Verified Alumni', headers, data_cse, widths)

    # 4. Recommended Demo Testing Flows
    p_g = doc.add_paragraph()
    r_g = p_g.add_run('4. Recommended Live Demonstration Workflows')
    r_g.font.name = 'Segoe UI'
    r_g.font.size = Pt(13)
    r_g.font.bold = True
    r_g.font.color.rgb = PRIMARY
    p_g.paragraph_format.space_before = Pt(12)
    p_g.paragraph_format.space_after = Pt(4)

    flows = [
        ('Faculty Posting Opportunities:', 'Login as faculty1.aiml@sbjit.edu.in -> Navigate to Opportunities -> Post academic research, internship, or job posting with required skills, stipend, and deadlines.'),
        ('Student Viewing & Applying to Faculty Postings:', 'Login as student1.aiml@sbjit.edu.in (or student1.cse) -> Navigate to Opportunities (/jobs) -> View Dr. Pradeep Mishra\'s postings in both the Explore feed (with Faculty badge) and the Faculty Opportunities tab -> Click "Apply to Faculty" -> Enter Statement of Interest -> Submit and track application in "My Applications".'),
        ('HOD Academic Dashboard:', 'Login as hod.aiml@sbjit.edu.in or hod.cse@sbjit.edu.in -> Navigate to Opportunities or Department -> View departmental participation stats, faculty KPI postings, and student collaboration metrics.'),
        ('TPO Placement & Candidate Review:', 'Login as tpo@sbjit.edu.in -> Manage verified company job postings, review student placement applications, and access the Student Finder.'),
        ('Alumni Referral Networking:', 'Login as alumni1.cse@sbjit.edu.in -> Post company openings or review student referral inquiries in the Alumni Directory.')
    ]

    for title, desc in flows:
        p_f = doc.add_paragraph()
        r_ft = p_f.add_run(f'• {title} ')
        r_ft.bold = True
        r_ft.font.color.rgb = DARK
        r_fd = p_f.add_run(desc)
        r_fd.font.color.rgb = SLATE
        p_f.paragraph_format.space_after = Pt(3)

    output_path = r'd:\yash\final_year_project\knots\KNOTS_Demo_Accounts_Credentials.docx'
    doc.save(output_path)
    print(f'Successfully generated: {output_path}')

if __name__ == '__main__':
    create_demo_doc()
